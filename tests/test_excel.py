"""Offline parity and adversarial OOXML tests; no database or UI dependencies."""
from copy import deepcopy
from io import BytesIO
import stat
import struct
from xml.etree import ElementTree as ET
from zipfile import ZIP_BZIP2, ZIP_DEFLATED, ZipFile, ZipInfo

from openpyxl import Workbook, load_workbook
import pytest

from employee_roster import excel
from employee_roster.domain import RosterError

SHEET = "xl/worksheets/sheet1.xml"
MAIN = excel.MAIN


def workbook_bytes(sheets=None):
    workbook = Workbook()
    workbook.remove(workbook.active)
    for name, rows in (sheets or [("Roster", [["DOMAINID", "EMAIL"], ["D001", "person@example.test"]])]):
        sheet = workbook.create_sheet(name)
        for row in rows:
            sheet.append(row)
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def parse(content, filename="roster.xlsx", mime=excel.XLSX_MIME):
    return excel.parse_workbook(content, filename, mime)


def rewrite(content, replacements=None, additions=None, remove=()):
    output = BytesIO()
    with ZipFile(BytesIO(content)) as source, ZipFile(output, "w", ZIP_DEFLATED) as target:
        for name in source.namelist():
            if name not in remove:
                target.writestr(name, (replacements or {}).get(name, source.read(name)))
        for name, value in (additions or {}).items():
            target.writestr(name, value)
    return output.getvalue()


def change_xml(content, path, change):
    with ZipFile(BytesIO(content)) as archive:
        root = ET.fromstring(archive.read(path))
    change(root)
    return rewrite(content, {path: ET.tostring(root)})


def cell(root, address):
    return root.find(f".//{MAIN}c[@r='{address}']")


def add_relationship(content, kind, target, mode=None, path="xl/_rels/workbook.xml.rels"):
    def change(root):
        attributes = {"Id": "securityTest", "Type": excel.DOC_REL_NS + "/" + kind, "Target": target}
        if mode is not None:
            attributes["TargetMode"] = mode
        ET.SubElement(root, f"{{{excel.PKG_REL_NS}}}Relationship", attributes)
    return change_xml(content, path, change)


def errors(result):
    return [issue for issue in result["issues"] if issue["severity"] == "error"]


@pytest.mark.parametrize("alias", ["DOMAINID", "DOMAIN", "Domain Id", "domain_id", " domain-id ", "D.o.m.a.i.n I.D.", "DOMAIN__ID", "DOMAIN_"])
def test_domain_aliases(alias):
    result = parse(workbook_bytes([("Roster", [[alias, "Email"], ["D001", "00123"]])]))
    assert result == {"rows": [{"row": 2, "domain": "D001", "values": {"EMAIL": "00123"}}],
                      "columns": ["DOMAINID", "EMAIL"], "issues": []}


def test_domain_can_be_after_identifier_columns():
    result = parse(workbook_bytes([("Roster", [["email", "domain"], ["001", "D001"]])]))
    assert result["columns"] == ["EMAIL", "DOMAINID"]
    assert result["rows"][0]["values"] == {"EMAIL": "001"}


def test_preserves_original_text_leading_zeros_unknown_columns_and_row_numbers():
    content = workbook_bytes([("Roster", [
        ["Domain Id", " email ", "Unknown-Code", "OTHER"],
        [" D001 ", " 00123 ", "Original Case", None],
        [None, None, None, None],
        ["d002", "00000", "0", "   "],
    ])])
    result = parse(content)
    assert result["rows"] == [
        {"row": 2, "domain": " D001 ", "values": {"EMAIL": " 00123 ", "UNKNOWN-CODE": "Original Case"}},
        {"row": 4, "domain": "d002", "values": {"EMAIL": "00000", "UNKNOWN-CODE": "0"}},
    ]
    assert result["columns"] == ["DOMAINID", "EMAIL", "UNKNOWN-CODE", "OTHER"]
    assert result["issues"] == []


def test_empty_cells_are_no_ops_and_missing_domain_is_skipped_with_warning():
    result = parse(workbook_bytes([("Roster", [
        ["DOMAINID", "EMAIL"], ["D001", " "], [None, "email"], [None, None], ["D002", None],
    ])]))
    assert result["rows"] == [{"row": 2, "domain": "D001", "values": {}}, {"row": 5, "domain": "D002", "values": {}}]
    assert len(result["issues"]) == 3
    assert next(issue for issue in result["issues"] if issue["row"] == 3)["column"] == "DOMAINID"
    assert all(issue["severity"] == "warning" for issue in result["issues"])


def test_blank_header_columns_are_ignored_without_shifting_cells():
    result = parse(workbook_bytes([("Roster", [["DOMAINID", None, "EMAIL"], ["D001", "ignored", "00123"]])]))
    assert result["columns"] == ["DOMAINID", "EMAIL"]
    assert result["rows"][0]["values"] == {"EMAIL": "00123"}


def test_duplicate_domains_are_retained_with_warning_for_later_validation():
    result = parse(workbook_bytes([("Roster", [["DOMAINID", "EMAIL"], ["D001", "one"], [" d001 ", "two"]])]))
    assert len(result["rows"]) == 2
    assert result["issues"] == [{"severity": "warning", "row": 3, "column": "DOMAINID",
                                 "message": "Duplicate DOMAINID in upload (also row 2)."}]


@pytest.mark.parametrize("headers", [["DOMAINID", "Email", " email "], ["Domain Id", "domain_id"], ["DOMAIN", "DOMAINID"]])
def test_duplicate_headers_are_errors_not_silent_overwrites(headers):
    result = parse(workbook_bytes([("Roster", [headers, ["D001", "first", "second"]])]))
    assert not result["rows"]
    assert any("Duplicate column" in issue["message"] and issue["row"] == 1 for issue in errors(result))


@pytest.mark.parametrize("rows", [[], [[None], ["data"]], [["NAME"], ["name"]], [["DOMAINID"]], [["DOMAINID", "EMAIL"], [None, "email"]]])
def test_empty_missing_header_missing_domain_and_no_importable_rows(rows):
    result = parse(workbook_bytes([("Roster", rows)]))
    assert not result["rows"]
    assert errors(result)


def test_roster_selection_is_case_insensitive_and_not_first_sheet():
    result = parse(workbook_bytes([
        ("Notes", [["DOMAINID", "EMAIL"], ["WRONG", "wrong"]]),
        ("rOsTeR", [["DOMAINID", "EMAIL"], ["RIGHT", "right"]]),
    ]))
    assert result["rows"][0]["domain"] == "RIGHT"
    assert not result["issues"]


def test_missing_roster_uses_first_sheet_and_warns():
    result = parse(workbook_bytes([("Employees", [["DOMAINID", "EMAIL"], ["D001", "mail"]])]))
    assert result["rows"][0]["domain"] == "D001"
    assert result["issues"][0]["row"] == 0
    assert "first worksheet 'Employees'" in result["issues"][0]["message"]


@pytest.mark.parametrize("invalid_rows", [[["WRONG"], ["data"]], [["DOMAINID"]], [["DOMAINID", "domain"], ["D001", "D001"]]])
def test_invalid_roster_falls_back_only_to_valid_first_sheet(invalid_rows):
    result = parse(workbook_bytes([
        ("Employees", [["DOMAINID", "EMAIL"], ["D001", "mail"]]), ("Roster", invalid_rows),
    ]))
    assert not errors(result)
    assert result["rows"][0]["domain"] == "D001"
    assert "falling back" in result["issues"][0]["message"]


def test_does_not_search_arbitrary_sheets_if_roster_and_first_are_invalid():
    result = parse(workbook_bytes([
        ("Notes", [["WRONG"]]), ("Roster", [["WRONG"]]), ("Third", [["DOMAINID"], ["D001"]]),
    ]))
    assert errors(result)
    assert not result["rows"]


@pytest.mark.parametrize("filename,mime", [
    ("file.xls", excel.XLSX_MIME), ("file.xlsm", excel.XLSX_MIME), ("file.xlsx.exe", excel.XLSX_MIME),
    ("file.xlsx", "application/zip"), ("file.xlsx", "application/octet-stream"), ("file.xlsx", ""),
    ("file.xlsx", "application/vnd.ms-excel"),
])
def test_requires_xlsx_extension_and_exact_mime(filename, mime):
    with pytest.raises(RosterError):
        parse(workbook_bytes(), filename, mime)


def test_extension_and_mime_case_are_insensitive():
    assert parse(workbook_bytes(), "ROSTER.XLSX", excel.XLSX_MIME.upper())["rows"]


@pytest.mark.parametrize("content", [b"", b"not excel", b"PK\x03\x04garbage", b"<workbook/>", b"\xd0\xcf\x11\xe0", None])
def test_invalid_bytes_raise_domain_error(content):
    with pytest.raises(RosterError):
        parse(content)


def test_generic_zip_is_not_an_xlsx_workbook():
    output = BytesIO()
    with ZipFile(output, "w", ZIP_DEFLATED) as archive:
        archive.writestr("readme.txt", "not an Excel workbook")
    with pytest.raises(RosterError, match="OOXML"):
        parse(output.getvalue())


@pytest.mark.parametrize("part", ["[Content_Types].xml", "_rels/.rels", "xl/workbook.xml", "xl/_rels/workbook.xml.rels", SHEET])
def test_missing_required_package_parts(part):
    with pytest.raises(RosterError):
        parse(rewrite(workbook_bytes(), remove=[part]))


@pytest.mark.parametrize("part", ["[Content_Types].xml", "xl/workbook.xml", SHEET, "docProps/core.xml"])
def test_malformed_xml_anywhere_is_rejected(part):
    with pytest.raises(RosterError):
        parse(rewrite(workbook_bytes(), {part: b"<broken"}))


@pytest.mark.parametrize("part", ["[Content_Types].xml", "xl/workbook.xml", SHEET, "docProps/core.xml"],
                         ids=["content-types", "workbook", "worksheet", "metadata"])
def test_unknown_xml_encoding_raises_domain_error(part):
    payload = b'<?xml version="1.0" encoding="not-a-real-encoding"?><root/>'
    with pytest.raises(RosterError, match="safely"):
        parse(rewrite(workbook_bytes(), {part: payload}))


def test_content_types_are_verified_not_just_zip_magic():
    def change(root):
        for element in root:
            if element.get("PartName") == "/xl/workbook.xml":
                element.set("ContentType", "application/xml")
    with pytest.raises(RosterError, match="XLSX"):
        parse(change_xml(workbook_bytes(), "[Content_Types].xml", change))


def test_workbook_without_sheets_is_rejected():
    def change(root):
        root.find(MAIN + "sheets").clear()
    with pytest.raises(RosterError, match="no worksheets"):
        parse(change_xml(workbook_bytes(), "xl/workbook.xml", change))


@pytest.mark.parametrize("payload", [
    b'<!DOCTYPE worksheet [<!ENTITY bomb "expanded">]><worksheet>&bomb;</worksheet>',
    b'<!DOCTYPE worksheet SYSTEM "file:///etc/passwd"><worksheet/>',
    b'<!DOCTYPE worksheet [<!ENTITY ext SYSTEM "https://example.test/data">]><worksheet>&ext;</worksheet>',
])
def test_dtd_and_entities_are_rejected(payload):
    with pytest.raises(RosterError, match="safely"):
        parse(rewrite(workbook_bytes(), {SHEET: payload}))


def test_deep_xml_is_rejected():
    payload = b"<root>" + b"<nested>" * 65 + b"</nested>" * 65 + b"</root>"
    with pytest.raises(RosterError, match="nesting"):
        parse(rewrite(workbook_bytes(), additions={"deep.xml": payload}))


@pytest.mark.parametrize("name", [
    "../evil.xml", "/absolute.xml", "a/../../evil.xml", "a\\evil.xml", "C:/evil.xml",
    "a//evil.xml", "a/./evil.xml", "%2e%2e/evil.xml", "evil.xml\x00hidden",
], ids=["parent", "absolute", "nested-parent", "backslash", "drive", "empty-segment",
        "dot-segment", "percent-encoded", "nul-byte"])
def test_unsafe_zip_paths(name):
    member = ZipInfo(name)
    # ZipInfo normalizes Windows separators and truncates NULs in its constructor.
    # Restore the malicious name so writestr emits it in BOTH ZIP headers.
    member.filename = name
    member.compress_type = ZIP_DEFLATED
    content = rewrite(workbook_bytes(), additions={member: b"<root/>"})
    with ZipFile(BytesIO(content)) as archive:
        stored = archive.infolist()[-1]
        assert stored.orig_filename == name
        offset = stored.header_offset
        assert content[offset:offset + 4] == b"PK\x03\x04"
        length = struct.unpack_from("<H", content, offset + 26)[0]
        assert content[offset + 30:offset + 30 + length] == name.encode("ascii")
    with pytest.raises(RosterError, match="ZIP member path"):
        parse(content)


def test_zip_symlink_is_rejected():
    member = ZipInfo("link.xml")
    member.create_system = 3
    member.external_attr = (stat.S_IFLNK | 0o777) << 16
    with pytest.raises(RosterError, match="member type"):
        parse(rewrite(workbook_bytes(), additions={member: b"target"}))


def test_duplicate_zip_member_is_rejected():
    with pytest.warns(UserWarning, match="Duplicate name"):
        content = rewrite(workbook_bytes(), additions={SHEET: b"<worksheet/>"})
    with pytest.raises(RosterError, match="Duplicate ZIP"):
        parse(content)


def test_case_colliding_zip_members_are_rejected():
    with pytest.raises(RosterError, match="Duplicate ZIP"):
        parse(rewrite(workbook_bytes(), additions={SHEET.upper(): b"<worksheet/>"}))


def test_unsupported_zip_compression_is_rejected():
    member = ZipInfo("extra.xml")
    member.compress_type = ZIP_BZIP2
    with pytest.raises(RosterError, match="compression"):
        parse(rewrite(workbook_bytes(), additions={member: b"<root/>"}))


def test_encrypted_flag_is_rejected_before_decompression():
    content = bytearray(workbook_bytes())
    central = content.index(b"PK\x01\x02")
    struct.pack_into("<H", content, central + 8, struct.unpack_from("<H", content, central + 8)[0] | 1)
    with pytest.raises(RosterError, match="Encrypted"):
        parse(bytes(content))


def test_corrupt_member_crc_becomes_domain_error():
    content = bytearray(workbook_bytes())
    central = content.index(b"PK\x01\x02")
    struct.pack_into("<I", content, central + 16, 0)
    with pytest.raises(RosterError, match="safely"):
        parse(bytes(content))


@pytest.mark.parametrize("target,mode", [("https://example.test/a.xlsx", "External"), ("file:///C:/secret.xlsx", "External"), ("https://example.test/a.xlsx", None), ("../../../secret.xml", None)])
def test_external_and_unsafe_relationships_are_rejected(target, mode):
    with pytest.raises(RosterError):
        parse(add_relationship(workbook_bytes(), "hyperlink", target, mode))


def test_external_hyperlink_on_unselected_sheet_is_rejected():
    workbook = Workbook()
    workbook.active.title = "Roster"
    workbook.active.append(["DOMAINID"])
    workbook.active.append(["D001"])
    other = workbook.create_sheet("Notes")
    other["A1"] = "Link"
    other["A1"].hyperlink = "https://example.test/"
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    with pytest.raises(RosterError, match="External"):
        parse(output.getvalue())


@pytest.mark.parametrize("part", ["xl/vbaProject.bin", "xl/externalLinks/externalLink1.xml", "xl/macrosheets/sheet1.xml", "xl/embeddings/object.bin"])
def test_macro_external_link_and_embedded_parts_are_rejected(part):
    with pytest.raises(RosterError):
        parse(rewrite(workbook_bytes(), additions={part: b"payload"}))


def test_renamed_macro_enabled_workbook_is_rejected():
    def change(root):
        for element in root:
            if element.get("PartName") == "/xl/workbook.xml":
                element.set("ContentType", "application/vnd.ms-excel.sheet.macroEnabled.main+xml")
    with pytest.raises(RosterError, match="macros"):
        parse(change_xml(workbook_bytes(), "[Content_Types].xml", change))


@pytest.mark.parametrize("sheet_name,hidden", [("Roster", False), ("Notes", False), ("Notes", True)])
def test_formulas_on_any_sheet_are_rejected(sheet_name, hidden):
    workbook = Workbook()
    workbook.active.title = "Roster"
    workbook.active.append(["DOMAINID", "EMAIL"])
    workbook.active.append(["D001", "mail"])
    sheet = workbook["Roster"] if sheet_name == "Roster" else workbook.create_sheet(sheet_name)
    sheet["B2"] = "=1+1"
    if hidden:
        sheet.sheet_state = "veryHidden"
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    with pytest.raises(RosterError, match="Formula"):
        parse(output.getvalue())


@pytest.mark.parametrize("attributes", [{}, {"t": "shared", "si": "0"}, {"t": "array", "ref": "B2:B3"}])
def test_cached_shared_and_array_formulas_are_rejected(attributes):
    def change(root):
        target = cell(root, "B2")
        target.clear()
        target.set("r", "B2")
        ET.SubElement(target, MAIN + "f", attributes).text = "1+1"
        ET.SubElement(target, MAIN + "v").text = "2"
    with pytest.raises(RosterError, match="Formula"):
        parse(change_xml(workbook_bytes(), SHEET, change))


@pytest.mark.parametrize("tag", ["formula", "formula1", "formula2", "calculatedColumnFormula", "totalsRowFormula"])
def test_non_cell_formula_elements_are_rejected(tag):
    def change(root):
        ET.SubElement(root, MAIN + tag).text = "1+1"
    with pytest.raises(RosterError, match="Formula"):
        parse(change_xml(workbook_bytes(), SHEET, change))


def test_formula_in_unreferenced_xml_part_is_rejected():
    with pytest.raises(RosterError, match="Formula"):
        parse(rewrite(workbook_bytes(), additions={"orphan.xml": b'<root><f>1+1</f></root>'}))


@pytest.mark.parametrize("name,value", [("Formula", "1+1"), ("_xlnm._FilterDatabase", "SUM(A1:A2)"), ("_xlnm._FilterDatabase", "'[other.xlsx]Roster'!A1:B2")])
def test_defined_name_formulas_and_disguised_external_ranges_are_rejected(name, value):
    def change(root):
        names = root.find(MAIN + "definedNames")
        if names is None:
            names = ET.SubElement(root, MAIN + "definedNames")
        ET.SubElement(names, MAIN + "definedName", {"name": name}).text = value
    with pytest.raises(RosterError, match="Formula"):
        parse(change_xml(workbook_bytes(), "xl/workbook.xml", change))


def test_limits_constants_are_production_bounds():
    assert excel.MAX_COMPRESSED_BYTES == 5 * 1024 * 1024
    assert excel.MAX_EXPANDED_BYTES == 50 * 1024 * 1024
    assert excel.MAX_DATA_ROWS == 10_000
    assert excel.MAX_COLUMNS == 100
    assert excel.MAX_CELLS == 500_000


def test_compressed_size_limit():
    with pytest.raises(RosterError, match="5 MB"):
        parse(b"PK\x03\x04" + b"x" * excel.MAX_COMPRESSED_BYTES)


def test_compressed_size_boundary(monkeypatch):
    content = workbook_bytes()
    monkeypatch.setattr(excel, "MAX_COMPRESSED_BYTES", len(content))
    assert parse(content)["rows"]
    monkeypatch.setattr(excel, "MAX_COMPRESSED_BYTES", len(content) - 1)
    with pytest.raises(RosterError, match="compressed"):
        parse(content)


def test_real_compressible_zip_bomb_is_rejected_before_member_read(monkeypatch):
    # Small test-scale ceiling exercises the same pre-decompression budget path.
    monkeypatch.setattr(excel, "MAX_EXPANDED_BYTES", 100_000)
    content = rewrite(workbook_bytes(), additions={"bomb.xml": b"x" * 200_000})
    monkeypatch.setattr(excel, "_read_part", lambda *args: pytest.fail("must reject before reading any member"))
    with pytest.raises(RosterError, match="expanded"):
        parse(content)


def test_claimed_50_mb_zip_bomb_is_rejected_without_allocation(monkeypatch):
    content = bytearray(workbook_bytes())
    central = content.index(b"PK\x01\x02")
    struct.pack_into("<I", content, central + 24, excel.MAX_EXPANDED_BYTES + 1)
    monkeypatch.setattr(excel, "_read_part", lambda *args: pytest.fail("must not decompress"))
    with pytest.raises(RosterError, match="50 MB"):
        parse(bytes(content))


def test_expanded_limit_is_aggregate_and_inclusive(monkeypatch):
    content = workbook_bytes()
    with ZipFile(BytesIO(content)) as archive:
        total = sum(member.file_size for member in archive.infolist())
    monkeypatch.setattr(excel, "MAX_EXPANDED_BYTES", total)
    assert parse(content)["rows"]
    monkeypatch.setattr(excel, "MAX_EXPANDED_BYTES", total - 1)
    with pytest.raises(RosterError, match="expanded"):
        parse(content)


def test_archive_entry_count_limit(monkeypatch):
    monkeypatch.setattr(excel, "MAX_ARCHIVE_ENTRIES", 1)
    with pytest.raises(RosterError, match="entries"):
        parse(workbook_bytes())


@pytest.mark.parametrize("dimension,reason", [("A1:XFD1048576", "row limit"), ("A1:CW2", "column limit"), ("A1:CV10001", "cell limit"), ("garbage", "reference"), ("B2:A1", "dimension")])
def test_dimension_bombs_and_malformed_dimensions(dimension, reason):
    def change(root):
        root.find(MAIN + "dimension").set("ref", dimension)
    with pytest.raises(RosterError, match=reason):
        parse(change_xml(workbook_bytes(), SHEET, change))


@pytest.mark.parametrize("address", ["A10002", "CW2", "XFD1048576", "A0", "A-1"])
def test_actual_cells_cannot_bypass_limits_with_small_dimensions(address):
    def change(root):
        root.find(MAIN + "dimension").set("ref", "A1")
        cell(root, "B2").set("r", address)
    with pytest.raises(RosterError):
        parse(change_xml(workbook_bytes(), SHEET, change))


def test_dimensions_are_checked_on_unselected_sheets():
    content = workbook_bytes([("Roster", [["DOMAINID"], ["D001"]]), ("Notes", [["note"]])])
    def change(root):
        root.find(MAIN + "dimension").set("ref", "A1:XFD1048576")
    with pytest.raises(RosterError):
        parse(change_xml(content, "xl/worksheets/sheet2.xml", change))


def test_last_allowed_row_and_column_are_accepted_without_dense_iteration():
    workbook = Workbook()
    workbook.active.title = "Roster"
    workbook.active["A1"] = "DOMAINID"
    workbook.active["A10001"] = "LAST"
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    assert parse(output.getvalue())["rows"][0]["row"] == 10001
    columns = ["DOMAINID", *[f"CODE{number}" for number in range(99)]]
    assert len(parse(workbook_bytes([("Roster", [columns, ["D001"]])]))["columns"]) == 100


def test_cell_limit_counts_header_and_all_sheets(monkeypatch):
    content = workbook_bytes()
    monkeypatch.setattr(excel, "MAX_CELLS", 4)
    assert parse(content)["rows"]
    monkeypatch.setattr(excel, "MAX_CELLS", 3)
    with pytest.raises(RosterError, match="cell limit"):
        parse(content)
    monkeypatch.setattr(excel, "MAX_CELLS", 7)
    with pytest.raises(RosterError, match="cell limit"):
        parse(workbook_bytes([("Roster", [["DOMAINID", "EMAIL"], ["D001", "mail"]]),
                              ("Other", [["DOMAINID", "EMAIL"], ["D002", "mail"]])]))


def test_duplicate_actual_cells_cannot_evade_cell_count(monkeypatch):
    monkeypatch.setattr(excel, "MAX_CELLS", 4)
    def change(root):
        row = root.find(f".//{MAIN}row[@r='2']")
        row.append(deepcopy(cell(root, "B2")))
    with pytest.raises(RosterError, match="cell limit"):
        parse(change_xml(workbook_bytes(), SHEET, change))


@pytest.mark.parametrize("change", [
    lambda root: root.find(f".//{MAIN}row[@r='2']").append(deepcopy(cell(root, "B2"))),
    lambda root: root.find(MAIN + "sheetData").append(deepcopy(root.find(f".//{MAIN}row[@r='2']"))),
    lambda root: cell(root, "B2").set("r", "B3"),
])
def test_duplicate_rows_cells_and_mismatched_coordinates_are_malformed(change):
    with pytest.raises(RosterError):
        parse(change_xml(workbook_bytes(), SHEET, change))


def test_shared_and_rich_text_preserve_zeros_and_ignore_phonetic_annotations():
    content = workbook_bytes()
    shared = (f'<sst xmlns="{excel.MAIN_NS}"><si><r><t>00</t></r><r><t>123</t></r>'
              '<rPh sb="0" eb="1"><t>annotation</t></rPh></si></sst>').encode()
    content = rewrite(content, additions={"xl/sharedStrings.xml": shared})
    content = add_relationship(content, "sharedStrings", "sharedStrings.xml")
    def content_type(root):
        ET.SubElement(root, f"{{{excel.CT_NS}}}Override", {"PartName": "/xl/sharedStrings.xml", "ContentType": excel.SHARED_TYPE})
    content = change_xml(content, "[Content_Types].xml", content_type)
    def change(root):
        target = cell(root, "B2")
        target.clear()
        target.attrib.update({"r": "B2", "t": "s"})
        ET.SubElement(target, MAIN + "v").text = "0"
    assert parse(change_xml(content, SHEET, change))["rows"][0]["values"] == {"EMAIL": "00123"}


def xstring_workbook(text, storage, address="B2", rich=False):
    """Inject wire-format ST_Xstring text, rather than relying on a writer's escaping."""
    content = workbook_bytes()
    if storage == "shared":
        root = ET.Element(MAIN + "sst")
        container = ET.SubElement(root, MAIN + "si")
        if rich:
            container = ET.SubElement(container, MAIN + "r")
        ET.SubElement(container, MAIN + "t").text = text
        content = rewrite(content, additions={"xl/sharedStrings.xml": ET.tostring(root)})
        content = add_relationship(content, "sharedStrings", "sharedStrings.xml")
        def content_type(root):
            ET.SubElement(root, f"{{{excel.CT_NS}}}Override", {
                "PartName": "/xl/sharedStrings.xml", "ContentType": excel.SHARED_TYPE})
        content = change_xml(content, "[Content_Types].xml", content_type)

    def change(root):
        target = cell(root, address)
        target.clear()
        target.attrib.update({"r": address, "t": {"shared": "s", "inline": "inlineStr", "str": "str"}[storage]})
        if storage == "shared":
            ET.SubElement(target, MAIN + "v").text = "0"
        elif storage == "str":
            ET.SubElement(target, MAIN + "v").text = text
        else:
            container = ET.SubElement(target, MAIN + "is")
            if rich:
                container = ET.SubElement(container, MAIN + "r")
            ET.SubElement(container, MAIN + "t").text = text
    return change_xml(content, SHEET, change)


@pytest.mark.parametrize("storage", ["shared", "inline", "str"])
@pytest.mark.parametrize("encoded,expected", [
    ("_x0041_", "A"), ("_x005F_x0041_", "_x0041_"),
    ("_x005F_x005F_x005F_x0041_", "_x005F_x0041_"),
    ("_x005f_x0000_", "_x0000_"), ("_x005F_xD800_", "_xD800_"),
    ("_xD83D__xDE00_", "\U0001f600"), ("_xd83d__xde00_", "\U0001f600"),
    ("before_x000D__x000A_after_x0009_tab", "before\r\nafter\ttab"),
    ("_x0030_00123", "000123"), ("_x003D_1+1", "=1+1"),
    ("_x0041__x0042_", "AB"), ("_x005F__x0041_", "_A"),
    ("x005F_ _xZZZZ_ _x041_ _X0041_", "x005F_ _xZZZZ_ _x041_ _X0041_"),
], ids=["character", "literal-escape", "nested-literal", "literal-nul", "literal-surrogate",
        "surrogate-pair", "lowercase-hex", "xml-whitespace", "leading-zeros", "formula-text",
        "adjacent-escapes", "underscore-and-character", "non-escapes"])
def test_excel_authored_xstrings_decode_once(storage, encoded, expected):
    result = parse(xstring_workbook(encoded, storage))
    assert result["rows"][0]["values"] == {"EMAIL": expected}
    assert not result["issues"]


@pytest.mark.parametrize("storage", ["shared", "inline"])
def test_rich_text_xstrings_decode_without_including_phonetics(storage):
    content = xstring_workbook("_x005F_x0041__xD83D__xDE00_", storage, rich=True)
    path = "xl/sharedStrings.xml" if storage == "shared" else SHEET
    def change(root):
        container = root.find(MAIN + "si") if storage == "shared" else cell(root, "B2").find(MAIN + "is")
        ET.SubElement(ET.SubElement(container, MAIN + "rPh"), MAIN + "t").text = "_x0000_"
        ET.SubElement(ET.SubElement(container, MAIN + "r"), MAIN + "t").text = "_x0042_"
    assert parse(change_xml(content, path, change))["rows"][0]["values"] == {"EMAIL": "_x0041_\U0001f600B"}


@pytest.mark.parametrize("storage", ["shared", "inline"])
def test_rich_text_does_not_decode_escape_tokens_created_by_concatenation(storage):
    content = xstring_workbook("_x00", storage, rich=True)
    path = "xl/sharedStrings.xml" if storage == "shared" else SHEET
    def change(root):
        container = root.find(MAIN + "si") if storage == "shared" else cell(root, "B2").find(MAIN + "is")
        ET.SubElement(ET.SubElement(container, MAIN + "r"), MAIN + "t").text = "41_"
    assert parse(change_xml(content, path, change))["rows"][0]["values"] == {"EMAIL": "_x0041_"}


@pytest.mark.parametrize("storage", ["shared", "inline", "str"])
@pytest.mark.parametrize("encoded", ["bad_x0000_text", "bad_x0001_text", "bad_x000B_text", "_xFFFF_", "_xFFFE_",
                                     "_xD800_", "_xDC00_", "_xD800__x0041_", "_xDC00__xD800_",
                                     "_xD800__xD800_", "_xD800_text_xDC00_"])
def test_xstrings_reject_invalid_decoded_controls_and_surrogates(storage, encoded):
    with pytest.raises(RosterError, match="control|surrogate"):
        parse(xstring_workbook(encoded, storage))


@pytest.mark.parametrize("storage", ["shared", "inline"])
def test_xstring_headers_decode_before_alias_normalization(storage):
    content = xstring_workbook("DOMAIN_x005F_ID", storage, address="A1")
    assert parse(content)["columns"] == ["DOMAINID", "EMAIL"]
    content = xstring_workbook("DOMAIN_x005F_ID", storage, address="B1")
    assert any("Duplicate column" in issue["message"] for issue in errors(parse(content)))


@pytest.mark.parametrize("index", ["0", "-1", "garbage", "999999999999999999999999"])
def test_invalid_shared_string_indices(index):
    def change(root):
        target = cell(root, "B2")
        target.clear()
        target.attrib.update({"r": "B2", "t": "s"})
        ET.SubElement(target, MAIN + "v").text = index
    with pytest.raises(RosterError, match="shared string"):
        parse(change_xml(workbook_bytes(), SHEET, change))


def test_numeric_lexical_values_are_not_float_coerced_or_display_formatted():
    def change(root):
        target = cell(root, "B2")
        target.clear()
        target.attrib.update({"r": "B2", "t": "n"})
        ET.SubElement(target, MAIN + "v").text = "000123"
    assert parse(change_xml(workbook_bytes(), SHEET, change))["rows"][0]["values"]["EMAIL"] == "000123"


def test_boolean_values_have_invariant_text():
    result = parse(workbook_bytes([("Roster", [["DOMAINID", "FLAG"], ["D001", True], ["D002", False]])]))
    assert [row["values"]["FLAG"] for row in result["rows"]] == ["True", "False"]


def test_export_round_trip_matches_csharp_warning_behavior_and_layout():
    rows = [{"domain": "D001", "values": {"EMAIL": "user@example.test", "ROLE": "Admin"}},
            {"domain": "D002", "values": {}}]
    original = deepcopy(rows)
    content = excel.export_workbook(rows, ["EMAIL", "ROLE"])
    result = parse(content)
    assert rows == original
    assert result["columns"] == ["DOMAINID", "EMAIL", "ROLE"]
    assert result["rows"] == [{"row": index, **row} for index, row in enumerate(rows, 2)]
    assert len(result["issues"]) == 1
    assert result["issues"][0]["severity"] == "warning"
    workbook = load_workbook(BytesIO(content))
    try:
        assert workbook.sheetnames == ["Roster"]
        assert workbook.active.freeze_panes == "A2"
        assert workbook.active.auto_filter.ref == "A1:C3"
        assert workbook.active["A1"].font.bold
        assert workbook.active.column_dimensions["A"].width == 25
    finally:
        workbook.close()


@pytest.mark.parametrize("value", ["=1+1", "+SUM(A1:A2)", "-1+2", "@SUM(A1:A2)", "\t=1+1", "000123", "#N/A", "https://example.test", "'literal", "two\nlines"])
def test_export_all_values_are_explicit_text_including_formula_looking_strings(value):
    content = excel.export_workbook([{"domain": value, "values": {"CODE": value}}], ["CODE"])
    workbook = load_workbook(BytesIO(content), data_only=False)
    try:
        for row in workbook.active:
            for target in row:
                assert target.data_type == "s"
                assert target.number_format == "@"
        assert workbook.active["A2"].value == value
        assert workbook.active["B2"].value == value
    finally:
        workbook.close()
    assert parse(content)["rows"][0] == {"row": 2, "domain": value, "values": {"CODE": value}}
    with ZipFile(BytesIO(content)) as archive:
        root = ET.fromstring(archive.read(SHEET))
    assert not root.findall(".//" + MAIN + "f")


@pytest.mark.parametrize("value,encoded", [
    ("_x0041_", "_x005F_x0041_"),
    ("_x005F_x0041_", "_x005F_x005F_x005F_x0041_"),
    ("_x0000_", "_x005F_x0000_"), ("_xD800_", "_x005F_xD800_"),
    ("_xD83D__xDE00_", "_x005F_xD83D__x005F_xDE00_"),
    ("_x0041_x0042_", "_x005F_x0041_x005F_x0042_"),
    ("000123_x004a_", "000123_x005F_x004a_"),
    ("=1+1_x0041_", "=1+1_x005F_x0041_"),
    ("_X0041_", "_X0041_"), ("\U0001f600", "\U0001f600"),
])
def test_export_escapes_literal_xstrings_as_excel_text_and_round_trips(value, encoded):
    content = excel.export_workbook([{"domain": value, "values": {"CODE": value}}], ["CODE"])
    with ZipFile(BytesIO(content)) as archive:
        root = ET.fromstring(archive.read(SHEET))
    for address in ("A2", "B2"):
        target = cell(root, address)
        assert target.get("t") == "inlineStr"
        assert target.findtext(MAIN + "is/" + MAIN + "t") == encoded
    assert not root.findall(".//" + MAIN + "f")
    assert parse(content)["rows"][0] == {"row": 2, "domain": value, "values": {"CODE": value}}


def test_export_does_not_truncate_escape_expansion_at_excel_text_limit():
    value = "_x0041_" * 4681
    assert len(value) == 32767
    content = excel.export_workbook([{"domain": "D001", "values": {"CODE": value}}], ["CODE"])
    with ZipFile(BytesIO(content)) as archive:
        root = ET.fromstring(archive.read(SHEET))
    assert cell(root, "B2").findtext(MAIN + "is/" + MAIN + "t") == "_x005F_x0041_" * 4681
    assert parse(content)["rows"][0]["values"]["CODE"] == value


def test_export_blank_cells_are_explicit_text_and_absent_from_parsed_values():
    content = excel.export_workbook([{"domain": "D001", "values": {}}], ["CODE"])
    with ZipFile(BytesIO(content)) as archive:
        root = ET.fromstring(archive.read(SHEET))
    assert cell(root, "B2").get("t") == "inlineStr"
    assert parse(content)["rows"][0]["values"] == {}


def test_export_deduplicates_codes_and_ignores_unknown_values():
    content = excel.export_workbook([{"domain": "D001", "values": {"email": "0001", "UNKNOWN": "not exported"}}],
                                    ["", "EMAIL", "email", " ROLE "])
    result = parse(content)
    assert result["columns"] == ["DOMAINID", "EMAIL", "ROLE"]
    assert result["rows"][0]["values"] == {"EMAIL": "0001"}


@pytest.mark.parametrize("code", ["DOMAINID", "domain", "Domain Id", "DOMAIN_ID", "DOMAIN__ID",
                                  "DOMAIN_", "D_O_M_A_I_N_I_D", "D.o.m.a.i.n I.D."])
def test_export_rejects_reserved_codes_instead_of_silently_dropping_identifiers(code):
    rows = [{"domain": "D001", "values": {code: "must not disappear"}}]
    before = deepcopy(rows)
    with pytest.raises(RosterError, match="reserved"):
        excel.export_workbook(rows, ["EMAIL", code])
    assert rows == before


@pytest.mark.parametrize("code", ["DOMAIN_ID2", "DOMAIN_NAME", "DOMAIN_IDENTIFIER", "MY_DOMAIN", "DOMAINID_EXTRA"])
def test_export_preserves_distinct_codes_near_domain_aliases(code):
    content = excel.export_workbook([{"domain": "D001", "values": {code: "00123"}}], [code])
    assert parse(content)["rows"][0]["values"] == {code: "00123"}


def test_export_formula_looking_headers_are_also_text():
    content = excel.export_workbook([{"domain": "D001", "values": {"=1+1": "safe"}}], ["=1+1"])
    assert parse(content)["columns"] == ["DOMAINID", "=1+1"]


def test_export_empty_roster_produces_header_only_template():
    content = excel.export_workbook([], ["CODE"])
    result = parse(content)
    assert result["columns"] == ["DOMAINID", "CODE"]
    assert errors(result)


@pytest.mark.parametrize("value", [
    "bad\x00text", "bad\x01text", "bad\ud800text", "bad\ufffftext", "x" * 32768, 123, None,
], ids=["nul-byte", "control-character", "unpaired-surrogate", "xml-noncharacter",
        "overlong-32768", "integer", "none"])
def test_export_rejects_invalid_nontext_or_overlong_values_without_truncation(value):
    with pytest.raises(RosterError):
        excel.export_workbook([{"domain": "D001", "values": {"CODE": value}}], ["CODE"])


@pytest.mark.parametrize("limit,value,rows,codes", [
    ("MAX_DATA_ROWS", 1, [{"domain": "D001", "values": {}}] * 2, []),
    ("MAX_COLUMNS", 1, [], ["CODE"]),
    ("MAX_CELLS", 3, [{"domain": "D001", "values": {}}], ["CODE"]),
    ("MAX_COMPRESSED_BYTES", 1, [], []),
    ("MAX_EXPANDED_BYTES", 1, [], []),
])
def test_export_limits(monkeypatch, limit, value, rows, codes):
    monkeypatch.setattr(excel, limit, value)
    with pytest.raises(RosterError, match="limit"):
        excel.export_workbook(rows, codes)


def test_cell_dimension_budget_at_production_boundary():
    def at_limit(root):
        root.find(MAIN + "dimension").set("ref", "A1:CV5000")
    assert parse(change_xml(workbook_bytes(), SHEET, at_limit))["rows"]
    def over_limit(root):
        root.find(MAIN + "dimension").set("ref", "A1:CV5001")
    with pytest.raises(RosterError, match="cell limit"):
        parse(change_xml(workbook_bytes(), SHEET, over_limit))


@pytest.mark.parametrize("kind,value", [("n", "NaN"), ("n", "not a number"), ("b", "2"), ("e", "#DIV/0!"), ("unknown", "value")])
def test_invalid_cell_types_and_numeric_values_are_rejected(kind, value):
    def change(root):
        target = cell(root, "B2")
        target.clear()
        target.attrib.update({"r": "B2", "t": kind})
        ET.SubElement(target, MAIN + "v").text = value
    with pytest.raises(RosterError):
        parse(change_xml(workbook_bytes(), SHEET, change))


def test_safety_failures_never_trigger_sheet_fallback():
    content = workbook_bytes([
        ("First", [["DOMAINID"], ["SAFE"]]),
        ("Roster", [["DOMAINID", "CODE"], ["UNSAFE", "=1+1"]]),
    ])
    with pytest.raises(RosterError, match="Formula"):
        parse(content)


def test_issue_schema_is_stable():
    result = parse(workbook_bytes([("Employees", [["DOMAINID", "CODE"], ["D001", None], [None, "orphan"]])]))
    for issue in result["issues"]:
        assert set(issue) == {"severity", "row", "column", "message"}
        assert issue["severity"] in ("warning", "error")
        assert type(issue["row"]) is int
        assert isinstance(issue["column"], str)
        assert isinstance(issue["message"], str)
