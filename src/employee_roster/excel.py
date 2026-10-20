"""Bounded, offline OOXML roster import and text-only export.

Unsafe/malformed packages raise RosterError. Roster content problems are returned
as issues (row 0 denotes a workbook/sheet issue). Columns include DOMAINID.
Text is preserved verbatim; whitespace-only cells are omitted, not deletions.
Numeric cells retain their stored lexical value, not Excel's display formatting.
"""
from __future__ import annotations

from io import BytesIO
import posixpath
import re
import stat
from urllib.parse import urlsplit
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile, ZIP_DEFLATED, ZIP_STORED, ZipFile
from zlib import error as ZlibError

from defusedxml import ElementTree as SafeET
from defusedxml.common import DefusedXmlException
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from .domain import RosterError, is_domain_alias, key, require

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
MAX_COMPRESSED_BYTES = 5 * 1024 * 1024
MAX_EXPANDED_BYTES = 50 * 1024 * 1024
MAX_DATA_ROWS = 10_000
MAX_COLUMNS = 100
MAX_CELLS = 500_000
MAX_ARCHIVE_ENTRIES = 2_048
MAIN_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
DOC_REL_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PKG_REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
CT_NS = "http://schemas.openxmlformats.org/package/2006/content-types"
WORKBOOK_TYPE = XLSX_MIME + ".main+xml"
WORKSHEET_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"
SHARED_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml"
MAIN = "{" + MAIN_NS + "}"
_XSTRING_ESCAPE = re.compile(r"_x([0-9A-Fa-f]{4})_")


def _xml(data: bytes):
    return SafeET.fromstring(data, forbid_dtd=True)


def _events(data: bytes):
    return SafeET.iterparse(BytesIO(data), events=("start", "end"), forbid_dtd=True)


def _part_name(name: str) -> bool:
    return bool(name) and not (
        name.startswith("/") or "\\" in name or ":" in name or "%" in name
        or any(ord(char) < 32 for char in name)
        or any(part in ("", ".", "..") for part in name.split("/"))
    )


def _coordinate(value: str) -> tuple[int, int]:
    match = re.fullmatch(r"([A-Z]{1,3})([1-9][0-9]{0,6})", value)
    require(match is not None, "Malformed worksheet cell reference.")
    column = 0
    for letter in match[1]:
        column = column * 26 + ord(letter) - ord("A") + 1
    row = int(match[2])
    require(row <= MAX_DATA_ROWS + 1, "Workbook exceeds the 10000 data row limit.")
    require(column <= MAX_COLUMNS, "Workbook exceeds the 100 column limit.")
    return row, column


def _range(value: str) -> tuple[int, int]:
    ends = value.split(":")
    require(len(ends) in (1, 2), "Malformed worksheet dimension.")
    first = _coordinate(ends[0])
    last = _coordinate(ends[-1])
    require(first[0] <= last[0] and first[1] <= last[1], "Malformed worksheet dimension.")
    return last


def _read_part(archive: ZipFile, name: str) -> bytes:
    with archive.open(name) as stream:
        data = stream.read(MAX_EXPANDED_BYTES + 1)
    require(len(data) <= MAX_EXPANDED_BYTES, "Workbook exceeds the 50 MB expanded archive limit.")
    require(len(data) == archive.getinfo(name).file_size, "Invalid ZIP member size.")
    return data


def _inspect_xml(data: bytes) -> tuple[int, int]:
    """Scan every XML part before sheet selection; never iterate claimed ranges."""
    cells = 0
    max_row = max_column = 0
    stack = []
    worksheet = False
    for event, element in _events(data):
        local = element.tag.rsplit("}", 1)[-1]
        if event == "start":
            if not stack:
                worksheet = element.tag == MAIN + "worksheet"
            stack.append(element)
            require(len(stack) <= 64, "XML nesting is too deep.")
            require(local not in {"f", "formula", "formula1", "formula2", "calculatedColumnFormula", "totalsRowFormula"},
                    "Formula-bearing workbooks are not allowed.")
            require(local not in {"externalLink", "externalReferences", "oleObject", "ddeLink"},
                    "External links and embedded objects are not allowed.")
            if worksheet and element.tag == MAIN + "dimension":
                row, column = _range(element.get("ref", ""))
                max_row, max_column = max(max_row, row), max(max_column, column)
            if worksheet and element.tag == MAIN + "row":
                reference = element.get("r", "")
                require(bool(re.fullmatch(r"[1-9][0-9]{0,6}", reference)), "Malformed worksheet row.")
                require(int(reference) <= MAX_DATA_ROWS + 1, "Workbook exceeds the 10000 data row limit.")
                max_row = max(max_row, int(reference))
            if worksheet and element.tag == MAIN + "c":
                row, column = _coordinate(element.get("r", ""))
                max_row, max_column = max(max_row, row), max(max_column, column)
                cells += 1
                require(cells <= MAX_CELLS, "Workbook exceeds the 500000 cell limit.")
        else:
            if local == "definedName":
                # Excel/openpyxl represent an autofilter as a reserved named range.
                reference = element.text or ""
                match = re.fullmatch(r"(?:'(?:[^'\[\]]|'')+'|[A-Za-z_][A-Za-z0-9_]*)!(\$?[A-Z]{1,3}\$?[1-9][0-9]{0,6}:\$?[A-Z]{1,3}\$?[1-9][0-9]{0,6})", reference)
                require(element.get("name") == "_xlnm._FilterDatabase" and match is not None,
                        "Formula-bearing defined names are not allowed.")
                _range(match[1].replace("$", ""))
            stack.pop()
            element.clear()
            if stack:
                stack[-1].remove(element)
    return cells, max_row * max_column


def _relationships(archive: ZipFile, name: str, names: set[str]) -> dict:
    root = _xml(_read_part(archive, name))
    require(root.tag == f"{{{PKG_REL_NS}}}Relationships", "Invalid OOXML relationships.")
    base = "" if name == "_rels/.rels" else posixpath.dirname(posixpath.dirname(name))
    result = {}
    for element in root:
        require(element.tag == f"{{{PKG_REL_NS}}}Relationship", "Invalid OOXML relationship.")
        identifier = element.get("Id", "")
        target = element.get("Target", "")
        kind = element.get("Type", "")
        require(identifier and identifier not in result, "Duplicate or missing relationship ID.")
        require(element.get("TargetMode", "Internal") == "Internal", "External links are not allowed.")
        parsed = urlsplit(target)
        require(target and not (parsed.scheme or parsed.netloc or parsed.query or parsed.fragment
                                or "\\" in target or "%" in target), "Unsafe relationship target.")
        resolved = posixpath.normpath(posixpath.join(base, target)) if not target.startswith("/") else target[1:]
        require(_part_name(resolved) and resolved in names, "Missing or unsafe relationship target.")
        require(not any(token in kind.lower() for token in ("external", "vba", "macro", "oleobject"))
                and not kind.lower().endswith("/package"),
                "External links, macros and embedded packages are not allowed.")
        result[identifier] = (kind, resolved)
    return result


def _inspect_package(archive: ZipFile) -> tuple[list[tuple[str, str]], str | None]:
    members = archive.infolist()
    require(len(members) <= MAX_ARCHIVE_ENTRIES, "Too many ZIP entries.")
    names = set()
    folded_names = set()
    expanded = 0
    for member in members:
        name = member.filename
        require(member.orig_filename == name and _part_name(name[:-1] if member.is_dir() else name), "Unsafe ZIP member path.")
        require(name.casefold() not in folded_names, "Duplicate ZIP member.")
        folded_names.add(name.casefold())
        names.add(name)
        mode = member.external_attr >> 16
        require(stat.S_IFMT(mode) in (0, stat.S_IFREG, stat.S_IFDIR), "Unsafe ZIP member type.")
        require(not member.flag_bits & 1, "Encrypted ZIP entries are not allowed.")
        require(member.compress_type in (ZIP_STORED, ZIP_DEFLATED), "Unsupported ZIP compression.")
        require(not any(token in name.lower() for token in ("vbaproject", "externallinks", "macrosheets", "embeddings/")),
                "External links and macros are not allowed.")
        expanded += member.file_size
        require(expanded <= MAX_EXPANDED_BYTES, "Workbook exceeds the 50 MB expanded archive limit.")
    require({"[Content_Types].xml", "_rels/.rels"} <= names, "File is not an OOXML workbook.")
    root = _xml(_read_part(archive, "[Content_Types].xml"))
    require(root.tag == f"{{{CT_NS}}}Types", "Invalid OOXML content types.")
    overrides, defaults = {}, {}
    for element in root:
        content_type = element.get("ContentType", "")
        require(not any(token in content_type.lower() for token in ("macro", "vba", "externallink", "oleobject")),
                "External links and macros are not allowed.")
        if element.tag == f"{{{CT_NS}}}Override":
            part = element.get("PartName", "")
            require(part.startswith("/") and _part_name(part[1:]) and part[1:] in names
                    and part[1:] not in overrides, "Invalid content type part.")
            overrides[part[1:]] = content_type
        elif element.tag == f"{{{CT_NS}}}Default":
            extension = element.get("Extension", "").lower()
            require(extension and extension not in defaults, "Duplicate content type default.")
            defaults[extension] = content_type
        else:
            raise RosterError("Invalid OOXML content type entry.")
    total_cells = total_area = 0
    relations = {}
    for member in members:
        if member.is_dir():
            continue
        name = member.filename
        data = _read_part(archive, name)
        content_type = overrides.get(name, defaults.get(name.rsplit(".", 1)[-1].lower(), ""))
        require(bool(content_type) or name == "[Content_Types].xml", "Missing OOXML content type.")
        if name.lower().endswith((".xml", ".rels")) or content_type.endswith("+xml") or content_type == "application/xml":
            cells, area = _inspect_xml(data)
            total_cells += cells
            total_area += area
            require(max(total_cells, total_area) <= MAX_CELLS, "Workbook exceeds the 500000 cell limit.")
        if name.lower().endswith(".rels"):
            relations[name] = _relationships(archive, name, names)
    office = [target for kind, target in relations["_rels/.rels"].values() if kind == DOC_REL_NS + "/officeDocument"]
    require(len(office) == 1 and overrides.get(office[0]) == WORKBOOK_TYPE, "File is not an XLSX workbook.")
    workbook_path = office[0]
    workbook = _xml(_read_part(archive, workbook_path))
    require(workbook.tag == MAIN + "workbook", "Invalid OOXML workbook.")
    rel_path = posixpath.join(posixpath.dirname(workbook_path), "_rels", posixpath.basename(workbook_path) + ".rels")
    workbook_relations = relations.get(rel_path, {})
    sheets = []
    seen_names, seen_targets = set(), set()
    for sheet in workbook.findall(MAIN + "sheets/" + MAIN + "sheet"):
        name = sheet.get("name", "")
        kind, target = workbook_relations.get(sheet.get(f"{{{DOC_REL_NS}}}id"), ("", ""))
        require(name and name.casefold() not in seen_names and target not in seen_targets, "Duplicate or missing worksheet.")
        require(kind == DOC_REL_NS + "/worksheet" and overrides.get(target) == WORKSHEET_TYPE, "Invalid worksheet relationship.")
        seen_names.add(name.casefold())
        seen_targets.add(target)
        sheets.append((name, target))
    require(bool(sheets), "Workbook contains no worksheets.")
    shared = [target for kind, target in workbook_relations.values() if kind == DOC_REL_NS + "/sharedStrings"]
    require(len(shared) <= 1, "Multiple shared string tables.")
    require(not shared or overrides.get(shared[0]) == SHARED_TYPE, "Invalid shared string table.")
    return sheets, shared[0] if shared else None


def _valid_xml_text(value: str) -> bool:
    return all(char in "\t\n\r" or 32 <= ord(char) <= 0xD7FF
               or 0xE000 <= ord(char) <= 0xFFFD or 0x10000 <= ord(char) <= 0x10FFFF for char in value)


def _decode_xstring(value: str) -> str:
    # One pass preserves _x005F_x0041_ as literal _x0041_, not A.
    decoded = _XSTRING_ESCAPE.sub(lambda match: chr(int(match[1], 16)), value)
    if decoded != value:
        try:
            decoded = decoded.encode("utf-16-le", errors="surrogatepass").decode("utf-16-le")
        except UnicodeError as exc:
            raise RosterError("Malformed UTF-16 surrogate in Excel text.") from exc
        require(_valid_xml_text(decoded), "Excel text contains invalid control or XML characters.")
    return decoded


def _text(element) -> str:
    # Decode each text node separately; exclude phonetic annotations.
    return "".join(_decode_xstring(node.text or "") for node in
                   element.findall(MAIN + "t") + element.findall(MAIN + "r/" + MAIN + "t"))


def _shared_strings(data: bytes) -> list[str]:
    result = []
    root = None
    for event, element in _events(data):
        if root is None:
            root = element
            require(root.tag == MAIN + "sst", "Invalid shared string table.")
        if event == "end" and element.tag == MAIN + "si":
            result.append(_text(element))
            require(len(result) <= MAX_CELLS, "Too many shared strings.")
            root.remove(element)
    return result


def _sheet_cells(data: bytes, shared: list[str]) -> dict[int, dict[int, str]]:
    result = {}
    stack = []
    current_row = 0
    for event, element in _events(data):
        if event == "start":
            if not stack:
                require(element.tag == MAIN + "worksheet", "Invalid worksheet XML.")
            stack.append(element)
            if element.tag == MAIN + "row":
                require(len(stack) == 3 and stack[-2].tag == MAIN + "sheetData",
                        "Malformed worksheet row placement.")
                current_row = int(element.get("r"))
                require(current_row not in result, "Duplicate worksheet row.")
                result[current_row] = {}
            continue
        if element.tag == MAIN + "c":
            row, column = _coordinate(element.get("r", ""))
            require(len(stack) >= 2 and stack[-2].tag == MAIN + "row" and row == current_row,
                    "Cell reference does not match its row.")
            require(column not in result[row], "Duplicate worksheet cell.")
            kind = element.get("t", "n")
            value = element.findtext(MAIN + "v", "")
            if kind == "inlineStr":
                inline = element.find(MAIN + "is")
                value = _text(inline) if inline is not None else ""
            elif kind == "s":
                require(bool(re.fullmatch(r"[0-9]{1,9}", value)), "Invalid shared string index.")
                index = int(value)
                require(index < len(shared), "Invalid shared string index.")
                value = shared[index]
            elif kind == "b":
                require(value in ("0", "1"), "Invalid boolean cell.")
                value = "True" if value == "1" else "False"
            else:
                require(kind in ("n", "str", "d"), "Unsupported or error-valued cell.")
                if kind == "str":
                    value = _decode_xstring(value)
                if kind == "n" and value:
                    require(re.fullmatch(r"[+-]?(?:[0-9]+(?:\.[0-9]*)?|\.[0-9]+)(?:[Ee][+-]?[0-9]+)?", value) is not None,
                            "Invalid numeric cell.")
            result[row][column] = value
            element.clear()
            stack[-2].remove(element)
        elif element.tag == MAIN + "row":
            element.clear()
            stack[-2].remove(element)
        stack.pop()
    return result


def _issue(severity: str, row: int, column: str, message: str) -> dict:
    return {"severity": severity, "row": row, "column": column, "message": message}


def _header(value: str) -> str:
    return "DOMAINID" if is_domain_alias(value) else key(value)


def _parse_sheet(cells: dict[int, dict[int, str]]) -> dict:
    issues, rows = [], []
    headers = {column: _header(value) for column, value in sorted(cells.get(1, {}).items()) if value.strip()}
    seen = set()
    for column, code in headers.items():
        if code in seen:
            issues.append(_issue("error", 1, code, f"Duplicate column '{code}' at {get_column_letter(column)}1."))
        seen.add(code)
    if "DOMAINID" not in seen:
        issues.append(_issue("error", 1, "DOMAINID", "Row 1 must contain a DOMAINID column."))
    result = {"rows": rows, "issues": issues, "columns": list(headers.values())}
    if issues:
        return result
    domain_column = next(column for column, code in headers.items() if code == "DOMAINID")
    seen_domains = {}
    for row_number, values in sorted(cells.items()):
        if row_number == 1:
            continue
        domain = values.get(domain_column, "")
        identifiers = {code: values[column] for column, code in headers.items()
                       if column != domain_column and values.get(column, "").strip()}
        if not domain.strip():
            if identifiers:
                issues.append(_issue("warning", row_number, "DOMAINID", "Row skipped because DOMAINID is blank."))
            continue
        if not identifiers:
            issues.append(_issue("warning", row_number, "", "Row has a DOMAINID but no identifiers were provided."))
        if key(domain) in seen_domains:
            issues.append(_issue("warning", row_number, "DOMAINID",
                                 f"Duplicate DOMAINID in upload (also row {seen_domains[key(domain)]})."))
        else:
            seen_domains[key(domain)] = row_number
        rows.append({"row": row_number, "domain": domain, "values": identifiers})
    if not rows:
        issues.append(_issue("error", 0, "DOMAINID", "No importable rows were found."))
    return result


def parse_workbook(content: bytes, filename: str, mime: str) -> dict:
    """Read XLSX bytes without Excel evaluation, network access, or extraction."""
    require(isinstance(filename, str) and filename.lower().endswith(".xlsx"), "Only .xlsx files are allowed.")
    require(isinstance(mime, str) and mime.strip().lower() == XLSX_MIME, "An XLSX MIME type is required.")
    require(isinstance(content, bytes) and bool(content), "Workbook content must be nonempty bytes.")
    require(len(content) <= MAX_COMPRESSED_BYTES, "Workbook exceeds the 5 MB compressed limit.")
    require(content.startswith(b"PK\x03\x04"), "File is not an XLSX ZIP archive.")
    try:
        with ZipFile(BytesIO(content)) as archive:
            sheets, shared_path = _inspect_package(archive)
            shared = _shared_strings(_read_part(archive, shared_path)) if shared_path else []
            # Parse every sheet, including hidden sheets, before allowing fallback.
            parsed = {path: _parse_sheet(_sheet_cells(_read_part(archive, path), shared)) for _, path in sheets}
            first = sheets[0]
            roster = next((sheet for sheet in sheets if sheet[0].casefold() == "roster"), None)
            selected = roster or first
            result = parsed[selected[1]]
            if roster and roster != first and any(issue["severity"] == "error" for issue in result["issues"]):
                fallback = parsed[first[1]]
                if not any(issue["severity"] == "error" for issue in fallback["issues"]):
                    fallback["issues"].insert(0, _issue("warning", 0, "", f"Worksheet '{roster[0]}' is invalid; falling back to first worksheet '{first[0]}'."))
                    return fallback
            if roster is None:
                result["issues"].insert(0, _issue("warning", 0, "", f"Roster worksheet not found; using first worksheet '{first[0]}'."))
            return result
    except RosterError:
        raise
    except (BadZipFile, DefusedXmlException, ParseError, ValueError, LookupError, TypeError, OSError, RuntimeError, NotImplementedError, EOFError, ZlibError) as exc:
        raise RosterError("The Excel file is malformed or cannot be read safely.") from exc


def export_workbook(rows: list[dict], codes: list[str]) -> bytes:
    """Create a text-only Roster sheet; codes are identifiers, not domain aliases."""
    columns = ["DOMAINID"]
    for code in codes:
        require(isinstance(code, str), "Identifier codes must be text.")
        require(not is_domain_alias(code), "Identifier code is reserved for the DOMAINID column.")
        normalized = _header(code)
        if normalized and normalized not in columns:
            columns.append(normalized)
    require(len(rows) <= MAX_DATA_ROWS, "Export exceeds the 10000 data row limit.")
    require(len(columns) <= MAX_COLUMNS, "Export exceeds the 100 column limit.")
    require((len(rows) + 1) * len(columns) <= MAX_CELLS, "Export exceeds the 500000 cell limit.")
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Roster"
    sheet.freeze_panes = "A2"
    sheet.column_dimensions["A"].width = 25

    def write_row(number: int, values: list[str]) -> None:
        for column, value in enumerate(values, 1):
            require(isinstance(value, str), "Export values must be text.")
            require(len(value) <= 32767 and _valid_xml_text(value),
                    "Export value contains invalid XML text or exceeds Excel's text limit.")
            cell = sheet.cell(number, column)
            # Assign the type explicitly AFTER the value: openpyxl infers '=' as a formula.
            cell.value = value
            cell.data_type = "s"
            cell.number_format = "@"
            if number == 1:
                cell.font = Font(bold=True)

    try:
        write_row(1, columns)
        for number, row in enumerate(rows, 2):
            values = {key(code): value for code, value in row["values"].items()}
            write_row(number, [row["domain"], *(values.get(code, "") for code in columns[1:])])
        sheet.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{len(rows) + 1}"
        for column in range(2, len(columns) + 1):
            sheet.column_dimensions[get_column_letter(column)].width = min(40, max(10, max(
                len(sheet.cell(row, column).value or "") for row in range(1, len(rows) + 2)) + 2))
        # openpyxl writes inline text verbatim. Escape literal ST_Xstring tokens
        # after validation/layout; the public setter would truncate expanded text.
        for row in sheet:
            for cell in row:
                cell._value = re.sub(r"_(?=x[0-9A-Fa-f]{4}_)", "_x005F_", cell.value)
        output = BytesIO()
        workbook.save(output)
        content = output.getvalue()
        require(len(content) <= MAX_COMPRESSED_BYTES, "Export exceeds the 5 MB compressed limit.")
        with ZipFile(BytesIO(content)) as archive:
            require(sum(member.file_size for member in archive.infolist()) <= MAX_EXPANDED_BYTES,
                    "Export exceeds the 50 MB expanded archive limit.")
        return content
    finally:
        workbook.close()
